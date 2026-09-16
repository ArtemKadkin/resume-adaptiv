from fastapi import FastAPI, HTTPException, File, UploadFile, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
import ollama
import io
import PyPDF2
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import requests
from bs4 import BeautifulSoup

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

def extract_text_from_file(file_bytes, filename):
    text = ""
    try:
        if filename.endswith('.pdf'):
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text: text += page_text + "\n"
        elif filename.endswith('.docx'):
            doc = Document(io.BytesIO(file_bytes))
            for para in doc.paragraphs: text += para.text + "\n"
        elif filename.endswith('.txt'):
            text = file_bytes.decode('utf-8')
        else:
            raise HTTPException(status_code=400, detail="Неподдерживаемый формат")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    return text

def create_professional_docx(text, photo_bytes=None):
    doc = Document()
    if photo_bytes:
        try:
            photo_stream = io.BytesIO(photo_bytes)
            doc.add_picture(photo_stream, width=Inches(1.0), height=Inches(1.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        except: pass
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    for line in text.split('\n'):
        line = line.strip()
        if not line: continue
        if line.startswith('#'):
            doc.add_heading(line.lstrip('# ').strip(), level=1)
        elif line.startswith('**') and line.endswith('**'):
            p = doc.add_paragraph()
            p.add_run(line.replace('**', '')).bold = True
        elif line.startswith('* ') or line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)
    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream

@app.get("/")
def read_root():
    return {"message": "OK"}

@app.post("/analyze")
async def analyze_resume(job_description: str = Form(...), resume_text: str = Form(None), file: UploadFile = File(None)):
    final_text = extract_text_from_file(await file.read(), file.filename) if file else resume_text
    if not final_text: raise HTTPException(status_code=400, detail="No resume")
    prompt = f"Эксперт ATS 2026. Проанализируй резюме и вакансию. РЕЗЮМЕ: {final_text} ВАКАНСИЯ: {job_description}. Markdown."
    try:
        response = ollama.chat(model='gemma4:26b', messages=[{'role': 'user', 'content': prompt}])
        return {"status": "success", "analysis": response['message']['content']}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/rewrite")
async def rewrite_resume(job_description: str = Form(...), resume_text: str = Form(None), file: UploadFile = File(None)):
    final_text = extract_text_from_file(await file.read(), file.filename) if file else resume_text
    if not final_text: raise HTTPException(status_code=400, detail="No resume")
    prompt = f"Перепиши резюме по стандартам 2026. ТОЛЬКО ТЕКСТ. РЕЗЮМЕ: {final_text} ВАКАНСИЯ: {job_description}"
    try:
        response = ollama.chat(model='gemma4:26b', messages=[{'role': 'user', 'content': prompt}])
        return {"status": "success", "analysis": response['message']['content'], "original_text": final_text}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/cover-letter")
async def generate_cover_letter(job_description: str = Form(...), resume_text: str = Form(None), file: UploadFile = File(None)):
    final_text = extract_text_from_file(await file.read(), file.filename) if file else resume_text
    prompt = f"Напиши сопроводительное письмо. РЕЗЮМЕ: {final_text} ВАКАНСИЯ: {job_description}"
    try:
        response = ollama.chat(model='gemma4:26b', messages=[{'role': 'user', 'content': prompt}])
        return {"status": "success", "analysis": response['message']['content']}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/export-docx")
async def export_docx(resume_content: str = Form(...), photo: UploadFile = File(None)):
    try:
        photo_bytes = await photo.read() if photo else None
        file_stream = create_professional_docx(resume_content, photo_bytes)
        return StreamingResponse(file_stream, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": "attachment; filename=resume_2026.docx"})
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/scrape")
async def scrape_job(url: str = Form(...)):
    try:
        headers = {'User-Agent': 'Mozilla/5.0'}
        res = requests.get(url, headers=headers, timeout=10)
        soup = BeautifulSoup(res.text, 'html.parser')
        for s in soup(["script", "style", "header", "footer", "nav"]): s.decompose()
        text = "\n".join([p.get_text().strip() for p in soup.find_all(['p', 'li']) if len(p.get_text().strip()) > 20])
        return {"status": "success", "scraped_text": text}
    except Exception as e:
        raise HTTPException(status_code=400, detail=str(e))
